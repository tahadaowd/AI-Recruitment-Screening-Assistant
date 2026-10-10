import os
import re
import base64
from pypdf import PdfReader
from docx import Document

SUPPORTED_EXTENSIONS = (".pdf", ".docx")

def clean_text(text):
    if not text:
        return ""
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()

def extract_pdf_text(path):
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                print(f"[Warning] PDF is encrypted and cannot be read: {path}")
                return ""

        pages_text = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            pages_text.append(page_text)

        full_text = "\n".join(pages_text)
        return clean_text(full_text)
    except Exception as e:
        print(f"[Warning] Failed to extract text from corrupted/unreadable PDF '{path}': {e}")
        return ""

def extract_docx_text(path):
    try:
        doc = Document(path)
        content_lines = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                content_lines.append(paragraph.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    content_lines.append(" | ".join(row_text))

        full_text = "\n".join(content_lines)
        return clean_text(full_text)
    except Exception as e:
        print(f"[Warning] Failed to extract text from corrupted/unreadable DOCX '{path}': {e}")
        return ""

def extract_text(path):
    if not os.path.exists(path):
        print(f"[Error] File not found: {path}")
        return ""

    if os.path.getsize(path) == 0:
        print(f"[Warning] File is empty (0 bytes): {path}")
        return ""

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        text = extract_pdf_text(path)
    elif ext == ".docx":
        text = extract_docx_text(path)
    else:
        print(f"[Warning] Unsupported file format: {ext}")
        return ""

    if not text.strip():
        print(f"[Warning] No readable text found in file '{path}' (may contain scanned images only).")

    return text

def save_attachment(service, message_id, attachment_info, target_dir="data/cvs"):
    filename = attachment_info.get("filename", "unnamed_cv")
    safe_filename = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', filename)
    unique_filename = f"{message_id}_{safe_filename}"

    os.makedirs(target_dir, exist_ok=True)
    file_path = os.path.join(target_dir, unique_filename)

    attachment_id = attachment_info.get("attachment_id")
    raw_data = attachment_info.get("data")

    if attachment_id:
        try:
            result = service.users().messages().attachments().get(
                userId="me",
                messageId=message_id,
                id=attachment_id
            ).execute()
            data_str = result.get("data", "")
        except Exception as e:
            print(f"[Error] Failed to download attachment {filename} from Gmail: {e}")
            return None
    elif raw_data:
        data_str = raw_data
    else:
        print(f"[Error] No attachmentId or inline data found for {filename}")
        return None

    try:
        binary_data = base64.urlsafe_b64decode(data_str)
        with open(file_path, "wb") as f:
            f.write(binary_data)
        return file_path
    except Exception as e:
        print(f"[Error] Failed to decode and save attachment {filename}: {e}")
        return None

def package_candidate_profile(email_data, file_path, cv_text):
    sender = email_data.get("sender", "")
    name_match = re.match(r"^\"?([^\"<]+)\"?\s*<.*>$", sender)
    candidate_name = name_match.group(1).strip() if name_match else sender

    return {
        "candidate_name": candidate_name,
        "sender_email": sender,
        "email_subject": email_data.get("subject", ""),
        "email_date": email_data.get("date", ""),
        "message_id": email_data.get("id", ""),
        "cv_file_path": file_path,
        "cv_filename": os.path.basename(file_path) if file_path else "",
        "raw_text": cv_text
    }